"""The metadata spreadsheet built from the results CSV and a local template."""
import datetime as dt
import os
from argparse import Namespace

import pytest
from openpyxl import load_workbook
from PIL import Image

from extractor.cli import _naming_for
from extractor.metadata import Template, digitized_date, write_xlsx
from extractor.output import append_rows

from .test_output import mk

TEMPLATE = {
    "name_template": "x_b{box}_f{folder}_{release:04d}",
    "list_name_template": "x_b{box}_f{folder}_list",
    "columns": [["Filename", "{file_name}.pdf"], ["Title", "{title}"], ["Kind", "Text"],
                ["Date", "{release_date}"], ["Made", "{digitized}"], ["Spare", ""], ["Spare", ""]],
}


def cells(path):
    return [[c.value for c in row] for row in load_workbook(path).active.iter_rows()]


def test_rows_follow_the_template_and_skip_unnamed_releases(tmp_path):
    named = mk("a.jpg", subject="Campus opens new library")
    named.file_name, named.digitized = "x_b1_f2_0001", "2024-05-06"
    untitled = mk("b.jpg", date="1971-03-04")
    untitled.file_name = "x_b1_f2_0002"
    csv_path = tmp_path / "r.csv"
    append_rows(csv_path, [named, untitled, mk("c.jpg")])  # c has no file name
    assert write_xlsx(csv_path, Template(TEMPLATE), tmp_path / "r.xlsx") == 2
    assert cells(tmp_path / "r.xlsx") == [
        ["Filename", "Title", "Kind", "Date", "Made", "Spare", "Spare"],
        ["x_b1_f2_0001.pdf", "Campus opens new library", "Text", "1970-11-02", "2024-05-06", None, None],
        ["x_b1_f2_0002.pdf", "Words here now please", "Text", "1971-03-04", None, None, None],
    ]
    assert (tmp_path / "r.xlsx").stat().st_mode & 0o077 == 0


def test_formula_like_text_stays_text(tmp_path):
    row = mk("a.jpg", subject="=HYPERLINK(1)")
    row.file_name = "n"
    append_rows(tmp_path / "r.csv", [row])
    write_xlsx(tmp_path / "r.csv", Template(TEMPLATE), tmp_path / "r.xlsx")
    cell = load_workbook(tmp_path / "r.xlsx").active["B2"]
    assert cell.value == "=HYPERLINK(1)" and cell.data_type == "s"


@pytest.mark.parametrize("data", [{}, {"columns": []}, {"columns": [["A"]]}, {"columns": [["A", "{nope}"]]},
                                  {"columns": [["A", "{title"]]}])
def test_bad_template_is_refused(data):
    with pytest.raises(ValueError):
        Template(data)


def test_template_names_the_pdfs_unless_a_name_is_typed(tmp_path):
    args = Namespace(box="1", folder="2", first_release=1, name_template=None, list_name_template=None,
                     template=Template(TEMPLATE))
    naming = _naming_for(tmp_path, args)
    assert naming.release_name(7) == "x_b1_f2_0007" and naming.list_name() == "x_b1_f2_list"
    args.name_template = "{box}-{folder}-{release}"
    assert _naming_for(tmp_path, args).release_name(7) == "1-2-7"


def test_digitized_date_comes_from_the_pdf_photo_or_file(tmp_path):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4\n1 0 obj <</Type/Catalog/Pages 2 0 R>> endobj\n"
                    b"2 0 obj <</Type/Pages/Kids[3 0 R]/Count 1>> endobj\n"
                    b"3 0 obj <</Type/Page/Parent 2 0 R/MediaBox[0 0 10 10]>> endobj\n"
                    b"4 0 obj <</CreationDate(D:20190304120000Z)>> endobj\n"
                    b"trailer <</Root 1 0 R/Info 4 0 R/Size 5>>\n%%EOF\n")
    photo = tmp_path / "b.jpg"
    exif = Image.Exif()
    exif[306] = "2023:07:08 10:11:12"
    Image.new("RGB", (4, 4)).save(photo, exif=exif)
    plain = tmp_path / "c.png"
    Image.new("RGB", (4, 4)).save(plain)
    os.utime(plain, (dt.datetime(2020, 1, 2, 12).timestamp(),) * 2)
    assert digitized_date(pdf) == "2019-03-04"
    assert digitized_date(photo) == "2023-07-08"
    if not hasattr(plain.stat(), "st_birthtime") and os.name != "nt":  # Linux keeps no creation date
        assert digitized_date(plain) == "2020-01-02"
    assert digitized_date(tmp_path / "missing.pdf") == ""

"""The cataloguing spreadsheet: one .xlsx row per named release, laid out by a local template.

The layout (column headings, what goes in each, and the PDF naming convention) belongs to the
archive, so it lives in a template file kept out of git (local/metadata_template.json by default):

    {"name_template": "...{box}...{folder}...{release:04d}",
     "list_name_template": "...{box}...{folder}..._list",
     "columns": [["Filename", "{file_name}.pdf"], ["Title", "{title}"], ["Type", "Text"], ["Notes", ""]]}

Columns are [heading, value] pairs, in order; headings may repeat. A value is fixed text or uses
the fields below, so a blank one is a column to fill in by hand. The xlsx is rebuilt from the
results CSV, so corrections made there carry over (python -m extractor xlsx results.csv).
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
from pathlib import Path

PROJECT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = PROJECT_DIR / "local" / "metadata_template.json"
# What a column value can use: every results CSV column, plus these.
EXTRA_FIELDS = ("title",)  # the list-page subject, else the first five words
_PDF_DATE = re.compile(r"(?:D:)?(\d{4})(\d{2})?(\d{2})?")


class Template:
    def __init__(self, data: dict, source: str = "template"):
        from .output import COLUMNS

        for key in ("name_template", "list_name_template"):
            if not isinstance(data.get(key) or "", str):
                raise ValueError(f"{source}: {key!r} must be text, like \"Box{{box}}_Folder{{folder}}_{{release:03d}}\".")
        self.name_template = data.get("name_template") or None
        self.list_name_template = data.get("list_name_template") or None
        cols = data.get("columns")
        if not isinstance(cols, list) or not cols or not all(
                isinstance(c, list) and len(c) == 2 and all(isinstance(x, str) for x in c) for c in cols):
            raise ValueError(f"{source}: 'columns' must be a list of [heading, value] pairs of text.")
        self.columns: list[tuple[str, str]] = [(h, v) for h, v in cols]
        self._blank = {c: "" for c in (*COLUMNS, *EXTRA_FIELDS)}
        for heading, value in self.columns:
            try:
                value.format(**self._blank)
            except (KeyError, IndexError, ValueError) as exc:
                raise ValueError(f"{source}: column {heading!r} has a bad value {value!r} ({exc}); "
                                 f"use fixed text or fields like {{file_name}}, {{title}}, {{release_date}}.")

    @classmethod
    def load(cls, path: Path) -> "Template":
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except OSError as exc:
            raise ValueError(f"Cannot read the metadata template {path}: {exc.strerror or exc}")
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError(f"{path} is not valid JSON: {exc}")
        if not isinstance(data, dict):
            raise ValueError(f"{path} must hold one JSON object.")
        return cls(data, str(path))

    def row(self, values: dict) -> list[str]:
        fields = dict(self._blank, **values)  # a CSV from an older version lacks the newer columns
        fields["title"] = fields.get("subject") or fields.get("first_five_words") or ""
        return [value.format(**fields) for _, value in self.columns]


def load_template(path: Path | None) -> Template | None:
    """The template at path; with no path, the default one if it exists."""
    path = path or (DEFAULT_TEMPLATE if DEFAULT_TEMPLATE.exists() else None)
    return Template.load(path) if path else None


def digitized_date(path: Path) -> str:
    """When the source file was made, as yyyy-mm-dd: a PDF's own creation date, a photo's EXIF
    date taken, else the file's creation date (its last change on Linux, which keeps none)."""
    try:
        if path.suffix.lower() == ".pdf":
            import pypdfium2 as pdfium

            doc = pdfium.PdfDocument(path)
            try:
                if m := _PDF_DATE.match(doc.get_metadata_dict().get("CreationDate", "").strip()):
                    return _iso(m[1], m[2], m[3]) or _file_date(path)
            finally:
                doc.close()
        else:
            from PIL import Image

            with Image.open(path) as img:
                exif = img.getexif()
                taken = exif.get_ifd(0x8769).get(36867) or exif.get(306)  # DateTimeOriginal, DateTime
            if m := re.match(r"(\d{4}):(\d{2}):(\d{2})", str(taken or "")):
                return _iso(*m.groups()) or _file_date(path)
    except Exception:  # an unreadable file still has a date on disk
        pass
    return _file_date(path)


def _iso(y: str, m: str | None, d: str | None) -> str:
    try:
        return dt.date(int(y), int(m or 1), int(d or 1)).isoformat()
    except ValueError:
        return ""


def _file_date(path: Path) -> str:
    try:
        st = path.stat()
    except OSError:
        return ""
    made = getattr(st, "st_birthtime", None) or (st.st_ctime if os.name == "nt" else st.st_mtime)
    return dt.date.fromtimestamp(made).isoformat()


def write_xlsx(csv_path: Path, template: Template, xlsx_path: Path) -> int:
    """Rebuild the workbook from the results CSV: one row per release that has a file name, in
    CSV order. Every cell is text, so dates stay yyyy-mm-dd and nothing is read as a formula.
    Returns the number of rows; a CSV it cannot use raises ValueError."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from .output import load_results, plain_cell, replace_private

    rows = [r for r in load_results(csv_path) if r["file_name"]]
    wb = Workbook()
    ws = wb.active
    for values in [[h for h, _ in template.columns]] + [
            template.row({k: plain_cell(v) for k, v in r.items()}) for r in rows]:
        ws.append(values)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.data_type, cell.number_format = "s", "@"
    ws.freeze_panes = "A2"
    replace_private(xlsx_path, wb.save, binary=True)
    return len(rows)
